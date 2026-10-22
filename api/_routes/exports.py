# ruff: noqa: E501  (column lists read better one per line)
"""Excel exports, for project managers and superadmins.

  GET /export/projects.xlsx   every project the reader runs (a superadmin: all), with every detail a
                              person would want, and a Maintenance sheet laid out like the project listing
  GET /export/people.xlsx     everyone with an account, their role, contact, groups and projects
  POST /export/table.xlsx     any list as its table shows it, filtered and sorted ({what, title, head, rows})

Only what people read: no internal ids, storage keys or GPS internals, and
never the NRIC digits. Column names follow the reader's language. Each
export is recorded in the audit log (who exported what, and how many rows).
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field

from _lib import checks as yearly
from _lib.account import ROLE_LABEL, Account
from _lib.db import fetch_all, transaction
from _lib.i18n import tr
from _lib.web import role
from _routes import maintenance as mt
from _routes import projects as proj

router = APIRouter()
pm_only = role("project_manager")
SG = ZoneInfo("Asia/Singapore")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SP_STATUS = {1: "Not submitted", 2: "Submitted", 3: "Approved"}


def yes_no(v: Any) -> str | None:
    return None if v is None else ("Yes" if v else "No")


def local(v: datetime | None) -> datetime | None:
    # Excel has no time zones: Singapore time, as people read it.
    return v.astimezone(SG).replace(tzinfo=None) if v else None


Sheet = tuple[str, list[Any], list[list[Any]]] | tuple[str, list[Any], list[list[Any]], list[Any]]


def workbook(sheets: list[Sheet], lang: str, *, as_sent: frozenset[str] = frozenset()) -> bytes:
    """Sheets of (title, heading row, rows[, a row above the headings]). Headings follow the reader's
    language, except the sheets named in as_sent (the client list's own layout, which the import reads back)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        title, head, rows = sheet[0], sheet[1], sheet[2]
        keep = title in as_sent

        def say(h: Any, keep: bool = keep) -> Any:
            return h if keep or not isinstance(h, str) else (tr(h, lang) or h)

        lead = sheet[3] if len(sheet) > 3 else None  # type: ignore[misc]
        ws = wb.create_sheet(say(title)[:31])
        top = 2 if lead is not None else 1
        if lead is not None:
            ws.append(lead)
        ws.append([say(h) for h in head])
        for r in rows:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
        for cell in ws[top]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0E7F53")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(cell.value, date):
                cell.number_format = "mmm yyyy"
        ws.freeze_panes = f"B{top + 1}"
        if rows:
            ws.auto_filter.ref = f"A{top}:{get_column_letter(len(head))}{len(rows) + top}"
        for i, h in enumerate(head, start=1):
            col = ws.column_dimensions[get_column_letter(i)]
            longest = max([len(str(say(h)))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None])
            col.width = min(max(10, longest + 2), 48)
            for row in ws.iter_rows(min_row=top + 1, min_col=i, max_col=i):
                c = row[0]
                if isinstance(c.value, datetime):
                    c.number_format = "yyyy-mm-dd hh:mm"
                elif isinstance(c.value, date):
                    c.number_format = "yyyy-mm-dd"
                elif isinstance(c.value, str) and "\n" in c.value:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _record(acct: Account, what: str, n: int) -> None:
    """The audit log keeps who exported what: the export is an act on personal data."""
    with transaction(acct.uid) as cur:
        cur.execute("insert into data_exports (what, row_count, exported_by) values (%s, %s, %s)", (what, n, acct.uid))


def _send(data: bytes, name: str) -> Response:
    stamp = datetime.now(SG).strftime("%d %b %Y")
    return Response(
        data,
        media_type=XLSX,
        headers={"Content-Disposition": f'attachment; filename="{name} {stamp}.xlsx"', "Cache-Control": "no-store"},
    )


# ------------------------------------------------------------------ projects

PROJECT_HEAD = [
    "S/N", "Project", "Status", "Address", "Postal Code", "Homeowner", "Homeowner contact no.", "Homeowner email",
    "Project manager", "Contractor", "Crew", "Created", "Start date", "Target end date", "Days running", "Progress (%)",
    "Current milestone", "Late", "Attention", "Electricity retailer", "Retailer contract end date", "SP application status",
    "Sales", "Waterproofing required", "Group chat created", "No. of Panels (est.)", "No. of Panels (actual)", "Panel Wp",
    "kWp DC", "Inverter", "Inverter serial number", "Inverter collected", "Inverter date", "Panel installation completed",
    "Scaffolding removed", "Scaffolding removal date", "SP submission date", "Grid connection done",
    "Scheduled grid connection date", "RCB replacement required", "RCB replacement date", "PVL received date",
    "Pre-inspection date", "SP appointment letter received", "Meter replacement date", "Turn On Date (SP inspection)",
    "FusionSolar app access", "Documents and photos", "Handover signed by", "Handover signed on", "Handed over on",
    "Handed over by", "PPA", "Maintenance Plan", "1P or 3P", "Next yearly check", "Yearly checks done",
    "Roof Access", "Urgent Maintenance",
]  # fmt: skip

MAINT_HEAD = [
    "S/N", "Address", "Postal Code", "PPA", "No. of Panels", "Panel Wp", "kWp DC", "1P or 3P", "Inverter",
    "Turn On Date", "Next yearly check", "Yearly checks done", "Maintenance Plan", "Roof Access",
    "Urgent Maintenance due to Poor Generation", "Project manager", "Homeowner", "Homeowner contact no.", "Project",
    "Source", "Notes",
]  # fmt: skip


def ppa_text(s: dict[str, Any] | None) -> str | None:
    if not s or not s.get("ppa_kind"):
        return None
    return "Value buy" if s["ppa_kind"] == "value_buy" else f"{s['ppa_years']} Years"


def plan_text(s: dict[str, Any] | None) -> str | None:
    if not s or not s.get("plan_years"):
        return None
    n = s["plan_years"]
    return f"{n} years excluding 1st year" if s["plan_excludes_first_year"] else f"Free for {n} years"


def phase_text(s: dict[str, Any] | None) -> str | None:
    return {1: "Single-phase", 3: "3-phase"}.get((s or {}).get("phase"))


def panel_wp(s: dict[str, Any]) -> Any:
    panels = s.get("panels") or []
    if len(panels) == 1:
        return panels[0].get("wp")
    return "\n".join(f"{p['count']} Nos. x {p.get('wp') or '?'} Wp" for p in panels) or None


def maintenance_rows(acct: Account) -> list[list[Any]]:
    where, params = mt._visible(acct)
    rows = fetch_all(f"{mt.SELECT} where {where} order by m.turned_on_on nulls last, m.address", params)  # type: ignore[arg-type]
    done = yearly.done([r["system_id"] for r in rows])
    out = []
    for n, s in enumerate(rows, start=1):
        panels = s["panels"] or []
        next_check, checks_done = yearly.summary(s, done.get(s["system_id"]))
        out.append([
            n, s["address"], s["postal_code"], ppa_text(s), sum(int(p.get("count") or 0) for p in panels) or None,
            panel_wp(s), s["kwp"], phase_text(s), "\n".join(s["inverters"] or []) or None, s["turned_on_on"],
            next_check, checks_done, plan_text(s),
            yes_no(s["roof_access"]), s["urgent_note"] or ("Yes" if s["urgent"] else None),
            s["pm_name"] or s["pm_email"], s["h_name"] or s["h_email"] or s["homeowner_name"], s["homeowner_contact_no"],
            s["project_name"], "Project listing" if s["import_ref"] else ("Project" if s["project_id"] else "Added"),
            s["notes"],
        ])  # fmt: skip
    return out


@router.get("/export/projects.xlsx")
def export_projects(acct: Account = Depends(pm_only)) -> Response:
    listed = {p["id"]: p for p in proj._load(acct)}
    ids = list(listed)
    raw = {
        r["project_id"]: r
        for r in fetch_all(
            "select p.*, r.name as retailer_name, h.email as h_email, cb.full_name as closer, "
            "s.signed_at, coalesce(s.signer_name, su.full_name) as signer, "
            "(select count(*) from project_files f where f.project_id = p.project_id) as files "
            "from projects p left join electricity_retailers r on r.retailer_id = p.electricity_retailer_id "
            "left join users h on h.uid = p.homeowner_id left join users cb on cb.uid = p.closed_by "
            "left join project_signatures s on s.project_id = p.project_id left join users su on su.uid = s.signed_by "
            "where p.project_id = any(%(ids)s)",
            {"ids": ids},
        )
    }
    systems = {
        s["project_id"]: s for s in fetch_all("select * from maintenance_systems where project_id = any(%s)", (ids,))
    }
    system_done = yearly.done([s["system_id"] for s in systems.values()])
    rows = []
    for n, pid in enumerate(sorted(ids, key=lambda i: listed[i]["createdAt"]), start=1):
        p, r, s = listed[pid], raw[pid], systems.get(pid)
        crew = ", ".join(t["name"] or "" for t in p["team"] if t["role"] in ("contractor", "epc_team"))
        count = r["panel_quantity_actual"] or r["panel_quantity_estimate"]
        kwp = round(count * r["panel_capacity"] / 1000, 3) if count and r["panel_capacity"] else None
        rows.append([
            n, r["name"], p["statusLabel"], r["address"], r["postal_code"], p["homeowner"]["name"],
            r["homeowner_contact_no"], r["h_email"], p["pm"]["name"], p["contractor"]["label"], crew or None,
            local(r["created_at"]), r["installation_start_date"], r["target_end_date"], p["daysElapsed"], p["progress"],
            "All milestones" if p["milestone"] == 3 else f"Milestone {p['currentMilestone']}",
            "Yes" if any(f["kind"] == "overdue" for f in p["flags"]) else "No",
            "\n".join(f["text"] for f in p["flags"]) or None, r["retailer_name"], r["retailer_contract_end_date"],
            SP_STATUS.get(r["sp_application_status"]), r["sales"], yes_no(r["waterproofing"]),
            yes_no(r["create_group_chat"]), r["panel_quantity_estimate"], r["panel_quantity_actual"], r["panel_capacity"],
            kwp, r["inverter_to_order"], r["inverter_serial_number"], yes_no(r["inverter_collected"]), r["inverter_date"],
            r["installation_end_date"], yes_no(r["scaffolding_removal"]), r["scaffolding_removal_date"],
            r["sp_submission_date"], yes_no(r["inverter_commission_grid_connection"]), r["commission_date"],
            yes_no(r["rcb_breaker_replacement"]), r["rcb_breaker_replacement_date"], r["pvl_received_date"],
            r["pre_inspection_date"], r["sp_appointment_letter_received_date"], r["meter_replacement_date"],
            r["sp_turn_on_inspection_date"], yes_no(r["fusion_solar_app_access"]), r["files"], r["signer"],
            local(r["signed_at"]), local(r["closed_at"]), r["closer"], ppa_text(s), plan_text(s), phase_text(s),
            *(yearly.summary(s, system_done.get(s["system_id"])) if s else (None, None)),
            yes_no(s["roof_access"]) if s else None, (s["urgent_note"] or ("Yes" if s["urgent"] else None)) if s else None,
        ])  # fmt: skip
    upkeep = maintenance_rows(acct)
    lang = (acct.user or {}).get("language") or "en"
    # First, the company's own client list, sheet for sheet as it was sent (and as the import reads it);
    # then the app's project details and maintenance, one row per project and per system.
    client = client_sheets(acct, list(listed.values()))
    data = workbook(
        [*client, ("Project details", PROJECT_HEAD, rows), ("Maintenance", MAINT_HEAD, upkeep)],
        lang,
        as_sent=frozenset(c[0] for c in client),
    )
    _record(acct, "projects", len(rows) + len(upkeep))
    return _send(data, "GetHomeApps projects")


# ------------------------------------------------------------------ the client list's own layout

# As "9 Solar Home_ Project List" has them (scripts/import_client_list.py reads the same).
CLIENT_HEAD = [
    "Project Code", "Salutation", "Homeowner", "Address", "Postal Code", "Mobile", "Email", "Signed Date",
    "Salesperson", "Referrer", "Solar Plan", "Rate (Per kWh)", "PPA", "EPC", "No. of Panels", "Panel Wp",
    "kWp DC", "1P or 3P", "Inverter", "Serial Number", "Dongle Serial No.", "Installation Start",
    "Installation End", "Turn On", "Turn On Date",
]  # fmt: skip
OUTPUT_HEAD = [
    "#",
    "Project Code",
    "Date",
    "Address",
    "PPA",
    "Panel No.",
    "Panel Wp",
    "System \nSize kWp",
    "Est. Output/ Mth",
]
BILLING_HEAD = ["#", "Project Code", "Date", "Address", "PPA", "Billing Period\n-owner side"]
FEEDBACK_HEAD = ["Project Code", "Homeowner", "Address", "Feedback"]
WARRANTY_HEAD = ["Huawei Inverter", "1P or 3P", "Price ex. GST", "Price incl. GST", "Warranty"]


def _lines(v: list[str] | None) -> str | None:
    return "\n".join(x for x in (v or []) if x) or None


def _months_back(n: int) -> list[date]:
    t = datetime.now(SG).date().replace(day=1)
    out = []
    for _ in range(n):
        out.insert(0, t)
        t = (t - timedelta(days=1)).replace(day=1)
    return out


def client_sheets(acct: Account, projects: list[dict[str, Any]]) -> list[Sheet]:
    """The client list as the company keeps it: a client per row (turned on: its maintenance system; not yet:
    its project), then the solar output and billing periods by month, feedback, and inverter warranties."""
    where, params = mt._visible(acct)
    systems = fetch_all(f"{mt.SELECT} where {where}", params)  # type: ignore[arg-type]
    open_ids = [p["id"] for p in projects if p["status"] != "maintenance"]
    raw = (
        fetch_all(
            "select p.*, h.email as h_email, h.full_name as h_name from projects p "
            "left join users h on h.uid = p.homeowner_id where p.project_id = any(%s)",
            (open_ids,),
        )
        if open_ids
        else []
    )
    label = {p["id"]: p["contractor"]["label"] for p in projects}
    clients: list[list[Any]] = []
    for s in systems:
        panels = s["panels"] or []
        clients.append([
            s["project_code"], s["salutation"], s["h_name"] or s["homeowner_name"], s["address"], s["postal_code"],
            s["homeowner_contact_no"], s["homeowner_email"] or s["h_email"], s["signed_on"], s["salesperson"],
            s["referrer"], s["solar_plan"], s["rate_per_kwh"], ppa_text(s), s["epc"],
            sum(int(p.get("count") or 0) for p in panels) or None, panel_wp(s), s["kwp"], phase_text(s),
            _lines(s["inverters"]), _lines(s["inverter_serials"]), _lines(s["dongle_serials"]),
            s["installation_start"], s["installation_end"], 1 if s["turned_on_on"] else None, s["turned_on_on"],
        ])  # fmt: skip
    for r in raw:
        count = r["panel_quantity_actual"] or r["panel_quantity_estimate"]
        contractor = label.get(r["project_id"])
        clients.append([
            r["project_code"], r["salutation"], r["h_name"] or r["homeowner_name"], r["address"], r["postal_code"],
            r["homeowner_contact_no"], r["homeowner_email"] or r["h_email"], r["signed_on"], r["sales"], r["referrer"],
            r["solar_plan"], r["rate_per_kwh"], ppa_text(r), None if contractor in (None, "—") else contractor, count,
            r["panel_capacity"], round(count * r["panel_capacity"] / 1000, 3) if count and r["panel_capacity"] else None,
            phase_text(r), r["inverter_to_order"], r["inverter_serial_number"], r["dongle_serial"],
            r["installation_start_date"], r["installation_end_date"], None, None,
        ])  # fmt: skip
    clients.sort(key=lambda c: (c[0] is None, str(c[0] or ""), str(c[3] or "")))

    months = _months_back(12)
    sids = [s["system_id"] for s in systems]
    output: dict[int, dict[date, Any]] = {}
    billing: dict[int, dict[date, str]] = {}
    if sids:
        for o in fetch_all("select system_id, month, kwh from system_output where system_id = any(%s)", (sids,)):
            output.setdefault(o["system_id"], {})[o["month"]] = o["kwh"]
        for b in fetch_all("select system_id, month, period from system_billing where system_id = any(%s)", (sids,)):
            billing.setdefault(b["system_id"], {})[b["month"]] = b["period"]
    on = sorted((s for s in systems if s["turned_on_on"]), key=lambda s: (s["turned_on_on"], s["address"]))
    out_rows, bill_rows = [], []
    for n, s in enumerate(on, start=1):
        panels = s["panels"] or []
        got = [output.get(s["system_id"], {}).get(m) for m in months]
        est = s["est_monthly_kwh"]
        filled = sum(1 for g in got if g is not None)
        out_rows.append([
            n, s["project_code"], s["turned_on_on"], s["address"], ppa_text(s),
            sum(int(p.get("count") or 0) for p in panels) or None, panel_wp(s), s["kwp"], est, *got,
            sum((g for g in got if g is not None), Decimal(0)) if filled else None,
            est * filled if est is not None and filled else None,
        ])  # fmt: skip
        bill_rows.append([
            n, s["project_code"], s["turned_on_on"], s["address"], ppa_text(s), s["billing_period"],
            *[billing.get(s["system_id"], {}).get(m) for m in months],
        ])  # fmt: skip

    codes = {s["system_id"]: s for s in systems}
    pcodes = {r["project_id"]: r for r in raw}
    feedback = []
    for f in fetch_all("select system_id, project_id, body from client_feedback order by created_at"):
        s, p = codes.get(f["system_id"]), pcodes.get(f["project_id"])
        if not s and not p:
            continue  # not one this reader looks after
        who = (s["h_name"] or s["homeowner_name"]) if s else (p["h_name"] or p["homeowner_name"])
        feedback.append([(s or p)["project_code"], who, (s or p)["address"], f["body"]])
    warranties = [
        [
            w["model"],
            phase_text(w),
            w["price_ex_gst"],
            w["price_incl_gst"],
            f"{w['warranty_years']} years" if w["warranty_years"] else None,
        ]
        for w in fetch_all("select * from inverter_models order by model")
    ]
    return [
        ("Clients List", CLIENT_HEAD, clients, [len(clients)]),
        ("Actual Solar Output (kWh)", [*OUTPUT_HEAD, *months, "Totally", "Est. Output"], out_rows),
        ("Billing Period", [*BILLING_HEAD, *[m.strftime("%b") for m in months]], bill_rows),
        ("Feedback", FEEDBACK_HEAD, feedback),
        ("Inverter Warranties", WARRANTY_HEAD, warranties),
    ]


# ------------------------------------------------------------------ people

PEOPLE_HEAD = [
    "S/N", "Name", "Role", "Email", "Mobile", "Mobile verified", "Account", "Signed in before", "Contractor groups",
    "Projects", "Address", "Postal Code", "Messages in", "Disables on", "Enables on", "Why disabled", "Invited",
    "Created",
]  # fmt: skip


@router.get("/export/people.xlsx")
def export_people(acct: Account = Depends(pm_only)) -> Response:
    users = fetch_all("select * from users order by active desc, user_type, full_name nulls last, email")
    groups: dict[int, list[str]] = {}
    for g in fetch_all(
        "select m.user_id, g.name from contractor_group_members m join contractor_groups g on g.group_id = m.group_id order by g.name"
    ):
        groups.setdefault(g["user_id"], []).append(g["name"])
    projects: dict[int, set[str]] = {}
    for p in fetch_all(
        """
        select p.project_manager_id as uid, p.name from projects p where p.project_manager_id is not null
        union select p.homeowner_id, p.name from projects p where p.homeowner_id is not null
        union select a.user_id, p.name from project_assignments a join projects p using (project_id)
        union select m.user_id, p.name from projects p join contractor_group_members m on m.group_id = p.contractor_group_id
        """
    ):
        projects.setdefault(p["uid"], set()).add(p["name"])
    rows = []
    for n, u in enumerate(users, start=1):
        rows.append([
            n, u["full_name"], ROLE_LABEL.get(u["user_type"], u["user_type"]), u["email"], u["contact_no"],
            "Yes" if u["mobile_verified_at"] else "No", "Active" if u["active"] else "Disabled",
            "Yes" if u["clerk_user_id"] else "No", ", ".join(groups.get(u["uid"], [])) or None,
            ", ".join(sorted(projects.get(u["uid"], set()))) or None, u["address"], u["postal_code"],
            "Chinese (Simplified)" if u["language"] == "zh" else "English", u["disable_on"], u["enable_on"],
            u["disabled_reason"], local(u["invited_at"]), local(u["created_at"]),
        ])  # fmt: skip
    lang = (acct.user or {}).get("language") or "en"
    data = workbook([("People", PEOPLE_HEAD, rows)], lang)
    _record(acct, "people", len(rows))
    return _send(data, "GetHomeApps people")


# ------------------------------------------------------------------ access review

ACCESS_HEAD = [
    "S/N", "Name", "Role", "Email", "Account", "Role since", "Last sign-in", "Two-step verification",
    "Projects", "Disables on", "Keep as is?", "Change to make",
]  # fmt: skip
REVIEW_EVERY = timedelta(days=91)


@router.get("/export/access-review.xlsx")
def export_access_review(acct: Account = Depends(pm_only)) -> Response:
    """Everyone who can sign in, what their role lets them see, when they last signed in and whether they
    use two-step verification, for a superadmin to review each quarter and sign off (ISO 27001 A.5.18,
    docs/iso/access-control.md). The last two come from Clerk; without it they're left blank."""
    if acct.role != "superadmin":
        raise HTTPException(403, "Only a superadmin reviews everyone's access.")
    from _lib import clerk

    try:
        facts = clerk.sign_in_facts()
    except Exception as exc:  # the review is still useful without Clerk's columns
        print(f"[access review] Clerk unavailable: {type(exc).__name__}: {exc}")
        facts = {}
    users = fetch_all(
        "select * from users where not is_demo order by active desc, user_type, full_name nulls last, email"
    )
    since = {
        r["uid"]: r["at"]
        for r in fetch_all(
            "select entity_id::bigint as uid, max(occurred_at) as at from audit_log "
            "where entity_table = 'users' and changes ? 'user_type' and entity_id ~ '^[0-9]+$' group by 1"
        )
    }
    counts: dict[int, int] = {}
    for r in fetch_all(
        """
        select uid, count(distinct project_id)::int as n from (
          select project_manager_id as uid, project_id from projects where project_manager_id is not null
          union select homeowner_id, project_id from projects where homeowner_id is not null
          union select user_id, project_id from project_assignments
          union select m.user_id, p.project_id from projects p join contractor_group_members m on m.group_id = p.contractor_group_id
        ) x group by uid
        """
    ):
        counts[r["uid"]] = r["n"]
    rows = []
    for n, u in enumerate(users, start=1):
        f = facts.get(u["clerk_user_id"] or "", {})
        last = f.get("last_sign_in_ms")
        rows.append([
            n, u["full_name"], ROLE_LABEL.get(u["user_type"], u["user_type"]), u["email"],
            "Active" if u["active"] else "Disabled", local(since.get(u["uid"]) or u["created_at"]),
            local(datetime.fromtimestamp(last / 1000, SG)) if last else ("Never" if u["clerk_user_id"] is None else None),
            ("On" if f["two_step"] else "Off") if f else None,
            "All" if u["user_type"] in ("superadmin", "finance") else counts.get(u["uid"], 0),
            u["disable_on"], None, None,
        ])  # fmt: skip
    today = datetime.now(SG).date()
    signoff = [
        ["Reviewed by", None],
        ["Their role", None],
        ["Date of the review", None],
        ["Accounts changed or switched off after it", None],
        ["Next review due", today + REVIEW_EVERY],
        [
            "Downloaded",
            f"{today.isoformat()} by {acct.user.get('full_name') or acct.user.get('email')}" if acct.user else None,
        ],
    ]
    lang = (acct.user or {}).get("language") or "en"
    data = workbook([("Access", ACCESS_HEAD, rows), ("Sign-off", ["Item", "Answer"], signoff)], lang)
    _record(acct, "access_review", len(rows))
    return _send(data, "GetHomeApps access review")


# ------------------------------------------------------------------ any table, as filtered


# What a table's Download may be recorded as, and who may download it (Admin: the lists it looks at).
TABLES = {
    "projects": ("project_manager", "finance"),
    "maintenance": ("project_manager", "finance"),
    "files": ("project_manager",),
    "invoices": ("finance",),
    "site_visits": ("project_manager",),
    "people": ("project_manager",),
    "solar": ("project_manager", "finance"),
    "solar_report": ("project_manager", "finance"),
}
DAY = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class TableIn(BaseModel):
    what: str
    title: str = Field(min_length=1, max_length=60)
    head: list[str] = Field(min_length=1, max_length=80)
    rows: list[list[str | float | int | None]] = Field(max_length=20000)


@router.post("/export/table.xlsx")
def export_table(body: TableIn, acct: Account = Depends(role("project_manager", "finance"))) -> Response:
    """A list exactly as its table shows it, filtered and sorted (components/data-table.tsx): the rows come
    from the reader's own screen, so it's only what they can already see. Recorded like every export."""
    who = TABLES.get(body.what)
    if who is None:
        raise HTTPException(400, "That list can't be downloaded.")
    if acct.role != "superadmin" and acct.role not in who:
        raise HTTPException(403, "You can't download this list.")
    if any(len(r) != len(body.head) for r in body.rows):
        raise HTTPException(400, "Each row needs a cell for every column.")
    rows = [[date.fromisoformat(v) if isinstance(v, str) and DAY.match(v) else v for v in r] for r in body.rows]
    lang = (acct.user or {}).get("language") or "en"
    # Excel refuses these in a sheet's name, and a name over 31 characters.
    title = re.sub(r"[\[\]:*?/\\]", " ", body.title).strip()[:31] or "Sheet"
    data = workbook([(title, body.head, rows)], lang)
    _record(acct, body.what, len(rows))
    return _send(data, f"GetHomeApps {tr(title, lang) or title}")
