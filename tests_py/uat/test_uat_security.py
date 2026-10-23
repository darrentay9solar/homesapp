# ruff: noqa: E501  (one rule per line reads best here)
"""UAT · the ISO 27001 changes (Oct 2026), person by person, written from the brief:

two-step     once 9 Solar Home turns it on (REQUIRE_STAFF_MFA=1), staff who signed in without the second
             step open only their own account (to set it up); with it, everything as before; homeowners
             aren't asked; off, nobody is
access       the quarterly access review is a superadmin's alone
health       the uptime check and the content security policy's reports need no sign-in, and say nothing
             about the app
"""

from __future__ import annotations

import pytest

from _lib.web import TWO_STEP_NEEDED
from conftest import bearer
from uat import spec

pytestmark = pytest.mark.uat

WITHOUT = {"fva": [3, -1]}  # a password only
WITH = {"fva": [3, 3]}  # and a code from an authenticator app
PATHS = ["/projects", "/maintenance", "/solar", "/calendar", "/alerts", "/me", "/me/solar"]


def _h(world, actor: str, **fva) -> dict[str, str]:
    return bearer(world.people[actor]["clerk_user_id"], **fva)


@pytest.fixture
def required(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("REQUIRE_STAFF_MFA", "1")


@pytest.mark.parametrize("path", PATHS)
@pytest.mark.parametrize("actor", spec.ACTORS)
def test_two_step_by_person(api, world, monkeypatch, actor, path) -> None:
    url = f"/api/py{path}"
    before = api.get(url, headers=_h(world, actor, **WITHOUT)).status_code  # off: as always
    monkeypatch.setenv("REQUIRE_STAFF_MFA", "1")
    without = api.get(url, headers=_h(world, actor, **WITHOUT))
    if spec.two_step_blocked(actor) and not path.startswith("/me"):
        assert without.status_code == 403 and without.json()["error"] == TWO_STEP_NEEDED, without.text
    else:
        assert without.status_code == before, without.text
    assert api.get(url, headers=_h(world, actor, **WITH)).status_code == before


@pytest.mark.parametrize("actor", spec.ACTORS)
def test_the_account_says_whether_two_step_is_needed(api, world, required, actor) -> None:
    r = api.get("/api/py/me", headers=_h(world, actor, **WITHOUT))
    if r.status_code == 200:
        assert r.json()["settings"]["twoStepNeeded"] is spec.two_step_blocked(actor)
        assert api.get("/api/py/me", headers=_h(world, actor, **WITH)).json()["settings"]["twoStepNeeded"] is False


@pytest.mark.parametrize("actor", spec.ACTORS)
def test_the_access_review_is_a_superadmins(api, world, actor) -> None:
    r = api.get("/api/py/export/access-review.xlsx", headers=world.h(actor))
    assert r.status_code == (200 if actor == "sa" else 403), r.text[:200]
    if r.status_code == 200:
        assert r.content[:2] == b"PK" and "spreadsheetml" in r.headers["content-type"]


def test_the_access_review_lists_everyone_and_who_is_switched_off(api, world) -> None:
    import io

    from openpyxl import load_workbook

    r = api.get("/api/py/export/access-review.xlsx", headers=world.h("sa"))
    rows = [[c.value for c in row] for row in load_workbook(io.BytesIO(r.content)).worksheets[0].iter_rows()]
    for actor in spec.ACTORS:
        row = next((x for x in rows if world.people[actor]["email"] in x), None)
        assert row is not None, actor
        assert ("Disabled" if actor == "inactive" else "Active") in row, actor


def test_the_health_check_needs_no_sign_in_and_says_nothing_else(api) -> None:
    r = api.get("/api/py/health")
    assert r.status_code == 200 and set(r.json()) == {"ok", "database", "ms"} and r.json()["ok"] is True


@pytest.mark.parametrize(
    "body",
    [
        {
            "csp-report": {
                "violated-directive": "img-src",
                "blocked-uri": "https://x.example/a.png?email=a@b.c",
                "document-uri": "https://app/x?y=1",
            }
        },
        [{"type": "csp-violation", "body": {"effectiveDirective": "script-src", "blockedURL": "inline"}}],
        "not json at all",
        {},
    ],
)
def test_the_policys_reports_need_no_sign_in(api, body, capsys) -> None:
    r = (
        api.post("/api/py/csp-report", json=body)
        if not isinstance(body, str)
        else api.post("/api/py/csp-report", content=body)
    )
    assert r.status_code == 204
    assert "email=" not in capsys.readouterr().out, "no query strings in the log"
