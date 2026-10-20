# ruff: noqa: E501, RUF001  (made-up rows in the list's own layout; it writes date ranges with en dashes)
"""The project list import (scripts/import_client_list.py) and the client record (migration 0038).

A small made-up workbook in the list's own layout goes through the real
import, inside a transaction that's rolled back. Then the API shows the
record to project managers only.
"""

from __future__ import annotations

import sys
import uuid
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import psycopg
import pytest
from psycopg.rows import dict_row
from test_project_work import _onemap, new_project, team  # noqa: F401

from conftest import TEST_OWNER_URL, bearer

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import import_client_list as icl

TAG = uuid.uuid4().hex[:4].upper()


def code(n: int) -> str:
    return f"9SH{TAG}{n:02d}"


def workbook(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clients List-Mia"
    ws.append([3])
    ws.append(icl.CLIENT_HEADER)
    ws.append([code(1), "Mr", "Test Owner One", "1 Pytest Lane", 999001, 91234001, "one@example.com", datetime(2026, 1, 5), "Sam Sales", "A Friend",
               "Smart Green", 0.189999999999998, "7 Years", "Pangxing Engineering", 20, 620, 12.4, "3-phase", "SUN2000-10KTL-MAP0", "SN1", "DG1",
               datetime(2026, 2, 1), datetime(2026, 2, 3), 1, datetime(2026, 3, 1)])  # fmt: skip
    ws.append([code(2), "Ms", "Test Owner Two", "2 Pytest Lane", 999002, "+65 9123 4002", "two@example.com", datetime(2026, 9, 1), "Sam Sales", None,
               "Smart Green\n-$3500", "-", "5 Years", "BTK", 16, 620, 9.92, "single phase", "UN2000-5KTL-L1 ", None, None,
               "21-Sept-26", "30-Sept-26", None, None])  # fmt: skip
    ws.append([code(3), None, "Test Owner Three", "3 Pytest Lane", 999003, None, None, datetime(2026, 10, 1), "Sam Sales", None,
               "Value Green", 0.19, "-", None, None, None, None, None, None, None, None, "first week Dec", None, None, None])  # fmt: skip
    o = wb.create_sheet("Actual Solar Output (kWh)-Mia")
    o.append(["#", "Project Code", "Date", "Address", "PPA", "Panel No.", "Panel Wp", "System Size kWp", "Est. Output/ Mth",
              datetime(2026, 6, 1), datetime(2026, 7, 2), "Aug", None, None, None, None, None, None, None, None, None, None, "Totally", "Est. Output"])  # fmt: skip
    o.append([1, code(1), None, "1 Pytest Lane", "7 Years", 20, 620, 12.4, 1300.5, 1200, "-", 1350.25, 999])
    b = wb.create_sheet("Billing Period")
    b.append(["#", "Project Code", "Date", "Address", "PPA", "Billing Period\n-owner side", "Aug", "Sep"])
    b.append(
        [
            1,
            None,
            datetime(2026, 3, 1),
            "1 Pytest Lane",
            "7 Years",
            "Full Calendar Month",
            "18 Jul – 17 Aug 2026",
            "Full Calendar Month",
        ]
    )
    f = wb.create_sheet("Feedback-Yichen")
    f.append(["Project Code", "Homeowner", "Address", "Feedback"])
    f.append([code(1), "Test Owner One", "1 Pytest Lane", "Very happy with the install."])
    f.append([code(2), "Test Owner Two", "2 Pytest Lane", "Waiting for the date."])
    w = wb.create_sheet("Inverter Warranties-Mia")
    w.append(["Huawei Inverter", "1P or 3P", "Price ex. GST", "Price incl. GST", "Warranty"])
    w.append([f"SUN2000-10KTL-MAP{TAG}", "3-phase", 2091, 2279.19, "10 Years"])
    w.append(["SUN2000-10KTL-MAP0", "3-phase", 2091, 2279.19, "10 Years"])
    path = tmp_path / "Project List-2026-10-14.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def owner():
    with psycopg.connect(TEST_OWNER_URL, row_factory=dict_row) as conn:
        yield conn
        conn.rollback()


def run(owner, path: Path) -> tuple[icl.Workbook, icl.Counts]:
    wb = icl.read(path, date(2026, 10, 14))
    cur = owner.cursor()
    cur.execute("select set_config('app.actor_uid', '', false)")
    return wb, icl.write(cur, wb, date(2026, 10, 14), geocode=False)


# ------------------------------------------------------------------ reading the list


def test_the_list_reads(tmp_path) -> None:
    wb = icl.read(workbook(tmp_path), date(2026, 10, 14))
    one, two, three = wb.clients
    assert one.turned_on == date(2026, 3, 1) and one.rate == Decimal("0.19") and one.ppa_years == 7 and one.phase == 3
    assert one.contact == "+65 9123 4001" and one.serials == ["SN1"] and one.dongles == ["DG1"]
    assert two.turned_on is None and two.start == date(2026, 9, 21) and two.end == date(2026, 9, 30)
    assert two.plan == "Smart Green -$3500" and two.rate is None and two.phase == 1
    assert two.inverters == ["SUN2000-5KTL-L1"], "a missing leading S and a trailing space"
    assert three.start is None and three.ppa_kind is None
    assert any("first week Dec" in n for n in wb.notes)
    assert wb.output[code(1)] == {
        date(2026, 6, 1): Decimal("1200"),
        date(2026, 8, 1): Decimal("1350.25"),
    }, "'-' isn't a figure"
    assert wb.estimate[code(1)] == Decimal("1300.5")
    assert wb.billing[code(1)] == (
        "Full Calendar Month",
        {date(2026, 8, 1): "18 Jul – 17 Aug 2026", date(2026, 9, 1): "Full Calendar Month"},
    ), "found by address"
    assert wb.feedback == {code(1): "Very happy with the install.", code(2): "Waiting for the date."}


def test_an_unknown_value_stops_it(tmp_path) -> None:
    path = workbook(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb["Clients List-Mia"]["M3"] = "Seven-ish"
    wb.save(path)
    with pytest.raises(icl.RowError, match=code(1)):
        icl.read(path, date(2026, 10, 14))


@pytest.mark.parametrize(
    ("a", "b"),
    [
        ("SUN2000-10KTL-LC0", "SUN2000-10K-LC0"),
        ("sun2000-12ktl-mb0", "SUN2000-12K-MB0"),
        ("SUN2000 - 30K-MC0", "SUN2000-30K-MC0"),
    ],
)
def test_one_inverter_many_spellings(a, b) -> None:
    assert icl.model_key(a) == icl.model_key(b.replace(" ", ""))


@pytest.mark.parametrize(
    ("d", "n", "out"),
    [
        (date(2026, 1, 31), 1, date(2026, 2, 28)),
        (date(2026, 3, 1), 6, date(2026, 9, 1)),
        (date(2025, 12, 15), 12, date(2026, 12, 15)),
    ],
)
def test_months_later(d, n, out) -> None:
    assert icl.add_months(d, n) == out


# ------------------------------------------------------------------ writing it


def test_the_import_writes_everything(owner, tmp_path) -> None:
    wb, k = run(owner, workbook(tmp_path))
    assert (k.systems_added, k.projects_added, k.output_months, k.billing_months, k.feedback) == (1, 2, 2, 2, 2)
    s = owner.execute("select * from maintenance_systems where project_code = %s", (code(1),)).fetchone()
    assert s["six_month_due"] == date(2026, 9, 1) and s["one_year_due"] == date(2027, 3, 1)
    assert (
        s["homeowner_name"] == "Test Owner One"
        and s["salesperson"] == "Sam Sales"
        and s["est_monthly_kwh"] == Decimal("1300.50")
    )
    assert s["billing_period"] == "Full Calendar Month" and s["inverter_serials"] == ["SN1"]
    p = owner.execute("select * from projects where project_code = %s", (code(2),)).fetchone()
    assert p["status"] == "in_progress" and p["project_manager_id"] is None, "waiting for a manager"
    assert p["installation_start_date"] == date(2026, 9, 21) and p["target_end_date"] == date(2026, 9, 30)
    assert p["ppa_years"] == 5 and p["contractor_text"] == "BTK" and p["record_note"] is None
    q = owner.execute("select * from projects where project_code = %s", (code(3),)).fetchone()
    assert q["record_note"] and "no installation dates" in q["record_note"], "dates it set itself are said"
    assert q["homeowner_contact_no"] == "Not in the list"
    fb = owner.execute(
        "select system_id, project_id from client_feedback where import_ref = %s", (f"feedback#{code(2)}",)
    ).fetchone()
    assert fb["project_id"] == p["project_id"] and fb["system_id"] is None


def test_running_it_again_changes_nothing_new(owner, tmp_path) -> None:
    path = workbook(tmp_path)
    run(owner, path)
    _, k = run(owner, path)
    assert (k.systems_added, k.projects_added) == (0, 0) and (k.systems_updated, k.projects_updated) == (1, 2)


def test_an_existing_system_is_matched_by_its_address(owner, tmp_path) -> None:
    owner.execute(
        "insert into maintenance_systems (import_ref, address, postal_code, turned_on_on, six_month_due, one_year_due, notes) "
        "values (%s, '1 Pytest Lane', '999001', '2026-03-01', '2026-09-01', '2027-03-01', 'Set in the app')",
        (f"pytest-{TAG}",),
    )
    _, k = run(owner, workbook(tmp_path))
    assert (k.systems_added, k.systems_updated) == (0, 1)
    s = owner.execute("select * from maintenance_systems where import_ref = %s", (f"pytest-{TAG}",)).fetchone()
    assert s["project_code"] == code(1) and s["notes"] == "Set in the app", "what the app set is kept"


def test_closing_a_project_keeps_its_record(owner, tmp_path) -> None:
    run(owner, workbook(tmp_path))
    p = owner.execute("select * from projects where project_code = %s", (code(2),)).fetchone()
    owner.execute("select maintenance_from_project(p) from projects p where project_id = %s", (p["project_id"],))
    s = owner.execute("select * from maintenance_systems where project_id = %s", (p["project_id"],)).fetchone()
    assert (
        s["project_code"] == code(2) and s["ppa_years"] == 5 and s["salesperson"] == "Sam Sales" and s["epc"] == "BTK"
    )
    fb = owner.execute("select system_id from client_feedback where project_id = %s", (p["project_id"],)).fetchone()
    assert fb["system_id"] == s["system_id"], "the feedback follows it to maintenance"


# ------------------------------------------------------------------ who sees the record


def h(u) -> dict[str, str]:
    return bearer(u["clerk_user_id"])


def test_managers_see_the_project_record_and_warranty(client, fx, team) -> None:  # noqa: F811
    pid = new_project(client, team)
    model = f"SUN2000-10KTL-MAP{TAG}"
    fx.conn.execute(
        "insert into inverter_models (model_key, model, phase, price_ex_gst, price_incl_gst, warranty_years) values (%s, %s, 3, 2091, 2279.19, 10) "
        "on conflict (model_key) do nothing",
        (icl.model_key(model), model),
    )
    try:
        fx.conn.execute(
            "update projects set project_code = %s, signed_on = '2026-01-05', solar_plan = 'Smart Green', rate_per_kwh = 0.19, "
            "inverter_to_order = %s where project_id = %s",
            (code(9), model.replace("KTL", "K"), pid),
        )
        r = client.get(f"/api/py/projects/{pid}/record", headers=h(team["pm"]))
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["record"]["projectCode"] == code(9) and d["record"]["ratePerKwh"] == 0.19
        assert (
            d["warranties"][0]["known"]
            and d["warranties"][0]["years"] == 10
            and d["warranties"][0]["priceInclGst"] == 2279.19
        )
        assert (
            client.get(f"/api/py/projects/{pid}/record", headers=h(team["ho"])).status_code == 403
        ), "not the homeowner"
        assert client.get(f"/api/py/projects/{pid}/record", headers=h(team["epc"])).status_code == 403, "not the crew"
        other = fx.user("project_manager", full_name="Another PM")
        assert (
            client.get(f"/api/py/projects/{pid}/record", headers=h(other)).status_code == 404
        ), "not another manager's project"
    finally:
        fx.conn.execute("delete from inverter_models where model_key = %s", (icl.model_key(model),))


def test_only_managers_keep_these_records(fx) -> None:
    from _lib.db import transaction

    ho = fx.user("homeowner")
    with pytest.raises(psycopg.errors.InsufficientPrivilege), transaction(ho["uid"]) as cur:
        cur.execute("insert into inverter_models (model_key, model) values ('X', 'X')")


# ------------------------------------------------------------------ duplicates stop it


def test_a_postal_code_twice_in_the_list_stops_it(owner, tmp_path) -> None:
    path = workbook(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb["Clients List-Mia"]["E4"] = 999001  # the second client at the first one's postal code
    wb.save(path)
    with pytest.raises(icl.RowError, match="same postal code"):
        run(owner, path)
    assert (
        owner.execute("select count(*) n from projects where project_code = %s", (code(2),)).fetchone()["n"] == 0
    ), "nothing written"


def test_a_postal_code_already_in_the_app_stops_it(owner, tmp_path) -> None:
    owner.execute(
        "insert into projects (name, address, postal_code, homeowner_name, homeowner_contact_no, installation_start_date, "
        "target_end_date, status) values ('Someone else', '9 Other Road', '999002', 'X', '+65 9000 0001', current_date, "
        "current_date + 21, 'draft')"
    )
    with pytest.raises(icl.RowError, match="already used by another project"):
        run(owner, workbook(tmp_path))


def test_flats_can_share_a_postal_code_when_said(owner, tmp_path) -> None:
    path = workbook(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb["Clients List-Mia"]["E4"] = 999001
    wb.save(path)
    parsed = icl.read(path, date(2026, 10, 14))
    cur = owner.cursor()
    cur.execute("select set_config('app.actor_uid', '', false)")
    k = icl.write(cur, parsed, date(2026, 10, 14), geocode=False, allow_shared_postal=True)
    assert k.projects_added == 2


def test_the_projects_download_reads_back_as_the_client_list(client, fx, team, tmp_path) -> None:  # noqa: F811
    """Downloaded, the client list is laid out as the company sent it, so the import reads it back."""
    pid = new_project(client, team)
    fx.conn.execute(
        "update projects set project_code = %s, salutation = 'Ms', solar_plan = 'Smart Green', rate_per_kwh = 0.19, "
        "ppa_kind = 'ppa', ppa_years = 7, phase = 3, panel_quantity_estimate = 20, panel_capacity = 620 where project_id = %s",
        (f"9SHT{TAG}", pid),
    )
    r = client.get("/api/py/export/projects.xlsx", headers=bearer(team["pm"]["clerk_user_id"]))
    assert r.status_code == 200, r.text
    path = tmp_path / "download.xlsx"
    path.write_bytes(r.content)
    wb = icl.read(path, date(2026, 10, 19))
    mine = next(c for c in wb.clients if c.code == f"9SHT{TAG}")
    assert (mine.salutation, mine.plan, float(mine.rate), mine.ppa_kind, mine.ppa_years) == ("Ms", "Smart Green", 0.19, "ppa", 7)
    assert (mine.panels, mine.wp, float(mine.kwp), mine.phase, mine.turned_on) == (20, 620, 12.4, 3, None)


# ------------------------------------------------------------------ importing it in the app


@pytest.fixture
def imported(fx):
    """Removes what an in-app import of the sample list wrote (it commits, unlike the owner tests)."""
    yield
    like = f"9SH{TAG}%"
    fx.conn.execute("delete from client_feedback where import_ref like %s", (f"feedback#{like}",))
    fx.conn.execute("delete from maintenance_systems where project_code like %s", (like,))
    fx.conn.execute("delete from projects where project_code like %s", (like,))
    fx.conn.execute("delete from inverter_models where model_key = %s", (icl.model_key(f"SUN2000-10KTL-MAP{TAG}"),))


def test_a_superadmin_checks_then_imports_the_list_in_the_app(client, fx, tmp_path, imported, monkeypatch) -> None:
    from _lib import onemap

    monkeypatch.setattr(onemap, "resolve", lambda _a, _p: (_ for _ in ()).throw(onemap.GeocodeError("offline", "unreachable")))
    boss = fx.user("superadmin")
    data = workbook(tmp_path).read_bytes()
    url = "/api/py/import/client-list"
    h = bearer(boss["clerk_user_id"])

    r = client.post(url, headers=h, content=data, params={"fileName": "Project List-2026-10-14.xlsx"})
    assert r.status_code == 200, r.text
    check = r.json()
    assert check["applied"] is False and check["listDate"] == "2026-10-14" and check["clients"] == 3 and check["turnedOn"] == 1
    assert check["counts"]["projectsAdded"] == 2 and check["counts"]["systemsAdded"] == 1
    none = fx.conn.execute("select count(*) as n from projects where project_code like %s", (f"9SH{TAG}%",)).fetchone()
    assert none["n"] == 0, "a check writes nothing"

    r = client.post(url, headers=h, content=data, params={"apply": 1, "fileName": "Project List-2026-10-14.xlsx"})
    assert r.status_code == 200, r.text
    assert r.json()["applied"] is True and r.json()["counts"]["projectsAdded"] == 2
    made = fx.conn.execute("select count(*) as n from projects where project_code like %s", (f"9SH{TAG}%",)).fetchone()
    assert made["n"] == 2
    actors = fx.conn.execute(
        "select distinct actor_uid from audit_log where entity_table = 'projects' and action = 'insert' "
        "and changes -> 'project_code' ->> 'to' like %s",
        (f"9SH{TAG}%",),
    ).fetchall()
    assert [a["actor_uid"] for a in actors] == [boss["uid"]], "written as the person importing"

    again = client.post(url, headers=h, content=data, params={"apply": 1, "fileName": "x-2026-10-14.xlsx"}).json()
    assert again["counts"]["projectsAdded"] == 0 and again["counts"]["projectsUpdated"] == 2, "a second import adds nothing"


@pytest.mark.parametrize(
    ("who", "body", "status", "says"),
    [
        ("project_manager", b"PK..", 403, None),
        ("superadmin", b"", 400, "Choose"),
        ("superadmin", b"not a workbook", 400, "Excel"),
        ("superadmin", b"PK\x03\x04broken", 400, "couldn't be read"),
    ],
)
def test_an_import_is_refused(client, fx, who, body, status, says) -> None:
    u = fx.user(who)
    r = client.post("/api/py/import/client-list", headers=bearer(u["clerk_user_id"]), content=body)
    assert r.status_code == status, r.text
    if says:
        assert says in r.json()["error"]


def test_a_list_with_a_repeated_postal_code_imports_nothing(client, fx, tmp_path, imported) -> None:
    path = workbook(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Clients List-Mia"]
    ws["E4"] = 999001  # client 2 at client 1's postal code
    wb.save(path)
    boss = fx.user("superadmin")
    r = client.post("/api/py/import/client-list", headers=bearer(boss["clerk_user_id"]), content=path.read_bytes(), params={"apply": 1})
    assert r.status_code == 400 and "Nothing imported" in r.json()["error"]
    n = fx.conn.execute("select count(*) as n from projects where project_code like %s", (f"9SH{TAG}%",)).fetchone()["n"]
    assert n == 0
