# ruff: noqa: E501  (one assertion per line reads best here)
"""Meetings and calendars, maintenance reminders, Excel exports (migration 0032).

arrange     invitations are asked, not held; accepting holds the time; nobody is double-booked
block       blocking out time, a window on each day or whole days, makes a person busy
change      a new time asks everyone again; cancelling tells everyone and is final
who         people invite only who they work with; the database refuses what the API refuses
reminders   each reminder for a check goes once: soon, on the day, overdue
exports     project managers download Excel; each export is in the audit log; nobody else can
"""

from __future__ import annotations

import io
import json
import uuid
from datetime import date, datetime, time, timedelta
from zoneinfo import ZoneInfo

import psycopg
import pytest

from _lib import maintenance_reminders as mr
from _lib.db import transaction
from conftest import bearer

SG = ZoneInfo("Asia/Singapore")


def h(u) -> dict[str, str]:
    return bearer(u["clerk_user_id"])


def slot(days: int, hour: int, minute: int = 0) -> str:
    """A Singapore time some days from now, as the screens send it."""
    d = datetime.now(SG).date() + timedelta(days=days)
    return datetime.combine(d, time(hour, minute)).isoformat(timespec="minutes")


@pytest.fixture
def pm(fx):
    return fx.user("project_manager", full_name="Meeting PM")


@pytest.fixture
def pm2(fx):
    return fx.user("project_manager", full_name="Other PM")


@pytest.fixture
def owner(fx):
    return fx.user("homeowner", full_name="Lonely Homeowner")


@pytest.fixture
def tidy(fx):
    """Meetings and blocks made by a test are removed after it (owner connection)."""
    yield
    if fx.uids:
        fx.conn.execute("delete from meetings where organiser_id = any(%s)", (fx.uids,))
        fx.conn.execute("delete from calendar_blocks where user_id = any(%s)", (fx.uids,))


def arrange(client, who, people, start, end, **kw):
    return client.post(
        "/api/py/meetings",
        headers=h(who),
        json={"title": "Site walk", "people": [p["uid"] for p in people], "startsAt": start, "endsAt": end, **kw},
    )


# ------------------------------------------------------------------ arrange


def test_an_invitation_is_asked_and_accepting_holds_the_time(client, pm, pm2, tidy):
    r = arrange(client, pm, [pm2], slot(3, 10), slot(3, 11), onlineLink="https://teams.microsoft.com/l/meetup-join/abc")
    assert r.status_code == 200, r.text
    m = r.json()["meeting"]
    assert m["status"] == "waiting" and m["canEdit"] and not m["canRespond"]
    assert {p["uid"]: p["response"] for p in m["people"]} == {pm["uid"]: "accepted", pm2["uid"]: "pending"}

    # Not answered yet: the invitee isn't busy, the organiser is.
    q = {"people": f"{pm['uid']},{pm2['uid']}", "from": slot(3, 0)[:10], "to": slot(3, 0)[:10]}
    busy = {p["uid"]: p["busy"] for p in client.get("/api/py/calendar/busy", headers=h(pm), params=q).json()["people"]}
    assert busy[pm2["uid"]] == [] and len(busy[pm["uid"]]) == 1

    cal = client.get("/api/py/calendar", headers=h(pm2), params={"from": slot(0, 0)[:10], "to": slot(6, 0)[:10]}).json()
    mine = next(x for x in cal["meetings"] if x["id"] == m["id"])
    assert mine["canRespond"] and mine["myResponse"] == "pending" and cal["toAnswer"] >= 1

    r = client.post(f"/api/py/meetings/{m['id']}/respond", headers=h(pm2), json={"response": "accepted"})
    assert r.status_code == 200 and r.json()["meeting"]["status"] == "confirmed"
    busy = {p["uid"]: p["busy"] for p in client.get("/api/py/calendar/busy", headers=h(pm), params=q).json()["people"]}
    assert len(busy[pm2["uid"]]) == 1


def test_both_sides_are_told(client, fx, pm, pm2, tidy):
    m = arrange(client, pm, [pm2], slot(3, 15), slot(3, 16)).json()["meeting"]
    client.post(f"/api/py/meetings/{m['id']}/respond", headers=h(pm2), json={"response": "declined"})
    asked = fx.conn.execute(
        "select link from notifications where recipient_uid = %s and kind = 'meeting_request'", (pm2["uid"],)
    ).fetchall()
    told = fx.conn.execute(
        "select title from notifications where recipient_uid = %s and kind = 'meeting_response'", (pm["uid"],)
    ).fetchall()
    assert [n["link"] for n in asked] == [f"/calendar?meeting={m['id']}"]
    assert [n["title"] for n in told] == ["Meeting declined · Site walk"]
    assert client.get(f"/api/py/meetings/{m['id']}", headers=h(pm)).json()["status"] == "declined"


def test_nobody_is_double_booked(client, fx, pm, pm2, tidy):
    m = arrange(client, pm, [pm2], slot(4, 14), slot(4, 15)).json()["meeting"]
    client.post(f"/api/py/meetings/{m['id']}/respond", headers=h(pm2), json={"response": "accepted"})
    boss = fx.user("superadmin", full_name="Booking Boss")
    r = arrange(client, boss, [pm2], slot(4, 14, 30), slot(4, 16))
    assert r.status_code == 409 and "Other PM" in r.json()["error"]
    # Next to it is fine.
    assert arrange(client, boss, [pm2], slot(4, 15), slot(4, 16)).status_code == 200


def test_times_are_on_the_half_hour_and_not_past(client, pm, pm2, tidy):
    assert arrange(client, pm, [pm2], slot(2, 9, 15), slot(2, 10)).status_code == 400
    assert arrange(client, pm, [pm2], slot(-2, 9), slot(-2, 10)).status_code == 400
    assert arrange(client, pm, [pm2], slot(2, 10), slot(2, 9)).status_code == 400
    assert arrange(client, pm, [pm2], slot(2, 10), slot(2, 11), onlineLink="teams.example/abc").status_code == 400
    assert arrange(client, pm, [], slot(2, 10), slot(2, 11)).status_code == 400


# ------------------------------------------------------------------ block


def test_blocking_out_days_makes_a_person_busy(client, pm, pm2, tidy):
    day = slot(5, 0)[:10]
    r = client.post("/api/py/calendar/blocks", headers=h(pm2), json={"fromDate": day, "note": "Leave"})
    assert r.status_code == 200, r.text
    assert arrange(client, pm, [pm2], slot(5, 10), slot(5, 11)).status_code == 409
    cal = client.get("/api/py/calendar", headers=h(pm2), params={"from": day, "to": day}).json()
    (b,) = cal["blocks"]
    assert b["note"] == "Leave"
    assert client.delete(f"/api/py/calendar/blocks/{b['id']}", headers=h(pm)).status_code == 404  # not theirs
    assert client.delete(f"/api/py/calendar/blocks/{b['id']}", headers=h(pm2)).status_code == 200
    assert arrange(client, pm, [pm2], slot(5, 10), slot(5, 11)).status_code == 200


def test_a_window_on_each_day_of_a_range(client, pm2, tidy):
    r = client.post(
        "/api/py/calendar/blocks",
        headers=h(pm2),
        json={"fromDate": slot(6, 0)[:10], "toDate": slot(8, 0)[:10], "startTime": "13:00", "endTime": "24:00"},
    )
    assert r.status_code == 200 and r.json()["message"] == "Blocked out 3 days."
    cal = client.get("/api/py/calendar", headers=h(pm2), params={"from": slot(6, 0)[:10], "to": slot(8, 0)[:10]}).json()
    assert [b["startsAt"][11:16] for b in cal["blocks"]] == ["13:00"] * 3
    assert (
        client.post(
            "/api/py/calendar/blocks",
            headers=h(pm2),
            json={"fromDate": slot(6, 0)[:10], "startTime": "13:15", "endTime": "14:00"},
        ).status_code
        == 400
    )


# ------------------------------------------------------------------ change


def test_a_new_time_asks_everyone_again_and_cancelling_is_final(client, fx, pm, pm2, tidy):
    m = arrange(client, pm, [pm2], slot(7, 9), slot(7, 10)).json()["meeting"]
    client.post(f"/api/py/meetings/{m['id']}/respond", headers=h(pm2), json={"response": "accepted"})
    assert client.patch(f"/api/py/meetings/{m['id']}", headers=h(pm2), json={"title": "Mine now"}).status_code == 403
    r = client.patch(
        f"/api/py/meetings/{m['id']}", headers=h(pm), json={"startsAt": slot(7, 11), "endsAt": slot(7, 12)}
    )
    assert r.status_code == 200 and r.json()["meeting"]["status"] == "waiting"
    n = fx.conn.execute(
        "select count(*) as n from notifications where recipient_uid = %s and kind = 'meeting_changed'", (pm2["uid"],)
    ).fetchone()["n"]
    assert n == 1
    assert client.post(f"/api/py/meetings/{m['id']}/cancel", headers=h(pm)).status_code == 200
    assert client.get(f"/api/py/meetings/{m['id']}", headers=h(pm2)).json()["status"] == "cancelled"
    assert (
        client.post(f"/api/py/meetings/{m['id']}/respond", headers=h(pm2), json={"response": "accepted"}).status_code
        == 409
    )
    assert client.patch(f"/api/py/meetings/{m['id']}", headers=h(pm), json={"title": "Back on"}).status_code == 403


# ------------------------------------------------------------------ who


def test_people_invite_only_who_they_work_with(client, fx, pm, owner, tidy):
    r = arrange(client, owner, [pm], slot(3, 10), slot(3, 11))
    assert r.status_code == 403
    assert (
        client.get(
            "/api/py/calendar/busy",
            headers=h(owner),
            params={"people": str(pm["uid"]), "from": slot(1, 0)[:10], "to": slot(2, 0)[:10]},
        ).status_code
        == 403
    )
    assert pm["uid"] not in {
        p["uid"] for p in client.get("/api/py/calendar/contacts", headers=h(owner)).json()["people"]
    }
    # A manager can reach anyone; then the homeowner sees that meeting, and nobody else's.
    m = arrange(client, pm, [owner], slot(3, 12), slot(3, 13)).json()["meeting"]
    assert client.get(f"/api/py/meetings/{m['id']}", headers=h(owner)).status_code == 200
    stranger = fx.user("homeowner")
    assert client.get(f"/api/py/meetings/{m['id']}", headers=h(stranger)).status_code == 404


def test_the_database_refuses_what_the_api_refuses(fx, pm, pm2, tidy, client):
    m = arrange(client, pm, [pm2], slot(9, 10), slot(9, 11)).json()["meeting"]
    with pytest.raises(psycopg.errors.InsufficientPrivilege), transaction(pm2["uid"]) as cur:
        cur.execute("update meetings set title = 'Hijacked' where meeting_id = %s", (m["id"],))
    with pytest.raises(psycopg.errors.InsufficientPrivilege), transaction(pm2["uid"]) as cur:
        cur.execute(
            "update meeting_people set response = 'accepted' where meeting_id = %s and user_id = %s",
            (m["id"], pm["uid"]),
        )
    with pytest.raises(psycopg.errors.InsufficientPrivilege), transaction(pm["uid"]) as cur:
        cur.execute("delete from meetings where meeting_id = %s", (m["id"],))
    with pytest.raises(psycopg.errors.CheckViolation), transaction(pm["uid"]) as cur:
        cur.execute(
            "update meetings set starts_at = starts_at + interval '10 minutes' where meeting_id = %s", (m["id"],)
        )
    # Accepting over time already blocked out is refused by the database itself.
    fx.conn.execute("insert into calendar_blocks (user_id, starts_at, ends_at) values (%s, %s, %s)",
                    (pm2["uid"], datetime.fromisoformat(slot(9, 10)).replace(tzinfo=SG), datetime.fromisoformat(slot(9, 12)).replace(tzinfo=SG)))  # fmt: skip
    with (
        pytest.raises(psycopg.errors.RaiseException, match="already taken in your calendar"),
        transaction(pm2["uid"]) as cur,
    ):
        cur.execute(
            "update meeting_people set response = 'accepted' where meeting_id = %s and user_id = %s",
            (m["id"], pm2["uid"]),
        )
    with pytest.raises(psycopg.errors.InsufficientPrivilege), transaction(pm["uid"]) as cur:
        cur.execute(
            "insert into calendar_blocks (user_id, starts_at, ends_at) values (%s, now()::date + 3, now()::date + 4)",
            (pm2["uid"],),
        )


def test_meetings_are_in_the_audit_log(client, fx, pm, pm2, tidy):
    m = arrange(client, pm, [pm2], slot(10, 10), slot(10, 11)).json()["meeting"]
    boss = fx.user("superadmin")
    log = client.get("/api/py/audit", headers=h(boss), params={"page": "meetings"}).json()
    mine = [e for e in log["entries"] if e["location"]["key"] == f"meeting:{m['id']}"]
    assert {e["summary"] for e in mine} >= {"Arranged a meeting", f"Invited {pm2['full_name']}"}
    assert all(e["lockedReason"] for e in mine)


# ------------------------------------------------------------------ reminders


@pytest.fixture
def system(fx, pm):
    tag = f"pytest-{uuid.uuid4().hex[:8]}"
    today = datetime.now(SG).date()
    sid = fx.conn.execute(
        "insert into maintenance_systems (import_ref, address, panels, turned_on_on, six_month_due, one_year_due, run_by) "
        "values (%s, '9 Reminder Road', %s, %s, %s, %s, %s) returning system_id",
        (
            tag,
            json.dumps([{"count": 10, "wp": 620}]),
            today - timedelta(days=170),
            today + timedelta(days=10),
            today + timedelta(days=195),
            pm["uid"],
        ),
    ).fetchone()["system_id"]
    yield sid
    fx.conn.execute("delete from maintenance_systems where system_id = %s", (sid,))


def test_stage_for():
    d = date(2026, 11, 11)
    assert mr.stage_for(d, date(2026, 10, 27)) is None
    assert mr.stage_for(d, date(2026, 10, 28)) == "soon"
    assert mr.stage_for(d, d) == "due"
    assert mr.stage_for(d, date(2026, 11, 12)) == "overdue"
    assert mr.stage_for(d, date(2027, 2, 10)) is None


def test_a_check_is_booked_through_the_calendar(client, fx, pm, owner, system, tidy):
    fx.conn.execute("update maintenance_systems set homeowner_id = %s where system_id = %s", (owner["uid"], system))
    # The office books with the homeowner; the homeowner with the manager looking after it.
    d = client.get(f"/api/py/maintenance/{system}", headers=h(pm)).json()
    assert d["canBook"] and d["bookWith"] == [owner["uid"]]
    contacts = {p["uid"] for p in client.get("/api/py/calendar/contacts", headers=h(owner)).json()["people"]}
    assert pm["uid"] in contacts

    r = client.post(
        "/api/py/meetings",
        headers=h(owner),
        json={
            "title": "6-month check",
            "people": [pm["uid"]],
            "startsAt": slot(8, 10),
            "endsAt": slot(8, 12),
            "place": "site",
            "systemId": system,
            "checkKey": "six_month",
        },
    )
    assert r.status_code == 200, r.text
    m = r.json()["meeting"]
    assert m["kind"] == "maintenance_check" and m["system"]["id"] == system and m["checkKey"] == "six_month"

    six = next(
        c
        for c in client.get(f"/api/py/maintenance/{system}", headers=h(pm)).json()["system"]["checks"]
        if c["key"] == "six_month"
    )
    assert six["booked"]["meetingId"] == m["id"] and six["booked"]["agreed"] is False
    client.post(f"/api/py/meetings/{m['id']}/respond", headers=h(pm), json={"response": "accepted"})
    six = next(
        c
        for c in client.get(f"/api/py/maintenance/{system}", headers=h(pm)).json()["system"]["checks"]
        if c["key"] == "six_month"
    )
    assert six["booked"]["agreed"] is True

    # Somebody else's system can't be booked.
    stranger = fx.user("homeowner")
    r = client.post(
        "/api/py/meetings",
        headers=h(stranger),
        json={
            "title": "x",
            "people": [pm["uid"]],
            "startsAt": slot(8, 14),
            "endsAt": slot(8, 15),
            "systemId": system,
            "checkKey": "one_year",
        },
    )
    assert r.status_code in (403, 404)
    # The audit log says what happened.
    boss = fx.user("superadmin")
    log = client.get("/api/py/audit", headers=h(boss), params={"page": "meetings"}).json()
    assert "Booked a maintenance check" in {
        e["summary"] for e in log["entries"] if e["location"]["key"] == f"meeting:{m['id']}"
    }


def test_each_reminder_goes_once(fx, pm, system):
    def mine() -> list[dict]:
        return fx.conn.execute(
            "select title, link from notifications where recipient_uid = %s and kind = 'maintenance_due' order by notification_id",
            (pm["uid"],),
        ).fetchall()

    mr.run()
    assert [n["title"] for n in mine()] == ["Maintenance check coming up · 9 Reminder Road"]
    assert mine()[0]["link"] == f"/maintenance/{system}"
    mr.run()
    assert len(mine()) == 1
    # On the day, once more; after it, once more.
    due = fx.conn.execute("select six_month_due from maintenance_systems where system_id = %s", (system,)).fetchone()[
        "six_month_due"
    ]
    mr.run(datetime.combine(due, time(9), SG))
    mr.run(datetime.combine(due + timedelta(days=3), time(9), SG))
    mr.run(datetime.combine(due + timedelta(days=4), time(9), SG))
    assert [n["title"].split(" · ")[0] for n in mine()] == [
        "Maintenance check coming up",
        "Maintenance check due today",
        "Maintenance check overdue",
    ]
    # Done: nothing more.
    fx.conn.execute("update maintenance_systems set six_month_done_on = %s where system_id = %s", (due, system))
    fx.conn.execute("delete from maintenance_reminders where system_id = %s", (system,))
    mr.run(datetime.combine(due + timedelta(days=5), time(9), SG))
    assert len(mine()) == 3


# ------------------------------------------------------------------ exports


def test_project_managers_export_to_excel_and_it_is_audited(client, fx, pm, owner):
    from openpyxl import load_workbook

    r = client.get("/api/py/export/projects.xlsx", headers=h(pm))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Projects", "Maintenance"]
    head = [c.value for c in wb["Projects"][1]]
    assert head[:3] == ["S/N", "Project", "Status"] and "Turn On Date (SP inspection)" in head

    r = client.get("/api/py/export/people.xlsx", headers=h(pm))
    assert r.status_code == 200
    head = [c.value for c in load_workbook(io.BytesIO(r.content))["People"][1]]
    assert "Email" in head and not any("NRIC" in (x or "") or x == "IC" for x in head)

    made = fx.conn.execute(
        "select what from data_exports where exported_by = %s order by export_id", (pm["uid"],)
    ).fetchall()
    assert [x["what"] for x in made] == ["projects", "people"]
    boss = fx.user("superadmin")
    log = client.get("/api/py/audit", headers=h(boss), params={"page": "exports"}).json()
    assert {e["summary"] for e in log["entries"] if e["actor"] and e["actor"]["uid"] == pm["uid"]} == {
        "Exported projects to Excel",
        "Exported people to Excel",
    }

    assert client.get("/api/py/export/projects.xlsx", headers=h(owner)).status_code == 403
    assert client.get("/api/py/export/people.xlsx", headers=h(owner)).status_code == 403
    with pytest.raises(psycopg.errors.InsufficientPrivilege), transaction(owner["uid"]) as cur:
        cur.execute("insert into data_exports (what, row_count, exported_by) values ('people', 1, %s)", (owner["uid"],))
    with pytest.raises(psycopg.errors.InsufficientPrivilege), transaction(pm["uid"]) as cur:
        cur.execute("delete from data_exports where exported_by = %s", (pm["uid"],))
