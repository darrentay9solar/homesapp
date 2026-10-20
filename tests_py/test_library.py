# ruff: noqa: E501  (one assertion per line reads best here)
"""My Files like a shared drive (api/_routes/files.py), files with no project yet, and any table to Excel.

library      everyone on a project sees all its files, foldered by milestone section; nobody else does
download     a file opens in the browser, or with ?download=1 is saved under its own name
unassigned   staff only: upload with no project, then assign to a project's slot under the slot's own
             rules (not before approval, photos only where photos go); removing it is audited
table        any list downloads as its table shows it, recorded as an export; only staff, only known lists
"""

from __future__ import annotations

import io

import pytest
from test_project_work import PNG, _onemap, approve_both, h, new_project, team, upload  # noqa: F401

PDF = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"


def unassigned(client, who, data=PDF, ctype="application/pdf", name="quote.pdf"):
    link = client.post("/api/py/unassigned-files/upload-link", headers=h(who), json={"contentType": ctype, "size": len(data), "fileName": name})
    if link.status_code != 200:
        return link
    put = client.put(link.json()["uploadUrl"], content=data, headers={"Content-Type": ctype})
    assert put.status_code == 200, put.text
    return client.post("/api/py/unassigned-files", headers=h(who), json={"key": link.json()["key"], "fileName": name, "note": " For Jalan Kayu "})


@pytest.fixture
def tidy(fx):
    yield
    fx.conn.execute("delete from unassigned_files where file_name like 'pytest-%'")


# ------------------------------------------------------------------ library


def test_everyone_on_a_project_sees_all_its_files(client, team) -> None:  # noqa: F811
    pid = new_project(client, team)
    approve_both(client, team, pid)
    r = upload(client, team["epc"], pid, "utility_bill", name="bill.png")
    assert r.status_code == 200, r.text
    for who in ("pm", "ho", "admin", "epc"):
        lib = client.get("/api/py/library", headers=h(team[who])).json()
        mine = [f for f in lib["files"] if f["projectId"] == pid]
        assert [f["name"] for f in mine] == ["bill.png"], who
        f = mine[0]
        assert f["section"] == "pre1" and f["sectionLabel"] == "Before Milestone 1" and f["categoryLabel"] == "Utility Bill"
        assert f["mine"] is (who == "epc") and f["uploader"]["uid"] == team["epc"]["uid"]
        assert pid in {p["id"] for p in lib["projects"]}
        assert lib["sections"][0] == {"key": "pre1", "label": "Before Milestone 1"}
    out = client.get("/api/py/library", headers=h(team["outsider"])).json()
    assert pid not in {f["projectId"] for f in out["files"]} and pid not in {p["id"] for p in out["projects"]}


def test_only_staff_see_unassigned(client, team) -> None:  # noqa: F811
    assert client.get("/api/py/library", headers=h(team["pm"])).json()["unassigned"] is not None
    for who in ("ho", "admin", "epc"):
        assert client.get("/api/py/library", headers=h(team[who])).json()["unassigned"] is None, who


# ------------------------------------------------------------------ download


def test_a_file_opens_or_downloads_under_its_name(client, team) -> None:  # noqa: F811
    pid = new_project(client, team)
    approve_both(client, team, pid)
    fid = upload(client, team["pm"], pid, "utility_bill", name="bill.png").json()["id"]
    for download, how in ((False, "inline"), (True, "attachment")):
        r = client.get(f"/api/py/files/{fid}", headers=h(team["ho"]), params={"download": download}, follow_redirects=False)
        assert r.status_code == 302
        got = client.get(r.headers["location"])
        assert got.status_code == 200 and got.content == PNG
        assert got.headers["content-disposition"].startswith(how) and "bill.png" in got.headers["content-disposition"]
    assert client.get(f"/api/py/files/{fid}", headers=h(team["outsider"]), follow_redirects=False).status_code in (403, 404)


# ------------------------------------------------------------------ unassigned


@pytest.mark.parametrize("who", ["ho", "admin", "epc"])
def test_only_staff_upload_without_a_project(client, team, who) -> None:  # noqa: F811
    r = client.post("/api/py/unassigned-files/upload-link", headers=h(team[who]), json={"contentType": "application/pdf", "size": 10})
    assert r.status_code == 403


def test_an_unassigned_file_is_assigned_under_the_slots_rules(client, fx, team, tidy) -> None:  # noqa: F811
    r = unassigned(client, team["pm"], name="pytest-quote.pdf")
    assert r.status_code == 200, r.text
    fid = r.json()["id"]
    lib = client.get("/api/py/library", headers=h(team["pm"])).json()
    row = next(u for u in lib["unassigned"] if u["id"] == fid)
    assert row["name"] == "pytest-quote.pdf" and row["note"] == "For Jalan Kayu" and row["kind"] == "document"
    d = client.get(f"/api/py/unassigned-files/{fid}", headers=h(team["pm"]), params={"download": 1}, follow_redirects=False)
    assert d.status_code == 302 and client.get(d.headers["location"]).headers["content-disposition"].startswith("attachment")

    pid = new_project(client, team)
    r = client.post(f"/api/py/unassigned-files/{fid}/assign", headers=h(team["pm"]), json={"projectId": pid, "category": "utility_bill"})
    assert r.status_code == 403 and "approved" in r.json()["error"], "a slot that isn't open yet isn't open to a file from Unassigned either"
    approve_both(client, team, pid)
    r = client.post(f"/api/py/unassigned-files/{fid}/assign", headers=h(team["pm"]), json={"projectId": pid, "category": "panel_pictures"})
    assert r.status_code == 400 and "photos" in r.json()["error"]
    r = client.post(f"/api/py/unassigned-files/{fid}/assign", headers=h(team["pm"]), json={"projectId": pid, "category": "nope"})
    assert r.status_code == 400
    r = client.post(f"/api/py/unassigned-files/{fid}/assign", headers=h(team["pm"]), json={"projectId": pid, "category": "utility_bill"})
    assert r.status_code == 200, r.text
    on = client.get("/api/py/library", headers=h(team["ho"])).json()["files"]
    moved = next(f for f in on if f["id"] == r.json()["id"])
    assert moved["projectId"] == pid and moved["category"] == "utility_bill" and moved["uploader"]["uid"] == team["pm"]["uid"]
    assert fid not in {u["id"] for u in client.get("/api/py/library", headers=h(team["pm"])).json()["unassigned"]}
    assert client.post(f"/api/py/unassigned-files/{fid}/assign", headers=h(team["pm"]), json={"projectId": pid, "category": "utility_bill"}).status_code == 404
    got = client.get(f"/api/py/files/{moved['id']}", headers=h(team["ho"]))
    assert got.status_code == 200 and got.content == PDF, "the stored copy is the one uploaded"


def test_an_unassigned_file_is_removed_and_audited(client, fx, team, tidy) -> None:  # noqa: F811
    fid = unassigned(client, team["pm"], name="pytest-old.pdf").json()["id"]
    assert client.delete(f"/api/py/unassigned-files/{fid}", headers=h(team["ho"])).status_code == 403
    r = client.delete(f"/api/py/unassigned-files/{fid}", headers=h(team["pm"]))
    assert r.status_code == 200 and "restored" in r.json()["message"]
    log = fx.conn.execute(
        "select action::text as action, actor_uid from audit_log where entity_table = 'unassigned_files' and entity_id = %s order by audit_id",
        (str(fid),),
    ).fetchall()
    assert [a["action"] for a in log] == ["insert", "delete"] and {a["actor_uid"] for a in log} == {team["pm"]["uid"]}


def test_an_upload_that_isnt_unassigned_is_refused(client, team) -> None:  # noqa: F811
    r = client.post("/api/py/unassigned-files", headers=h(team["pm"]), json={"key": "projects/1/documents/utility_bill/x.pdf"})
    assert r.status_code == 400


# ------------------------------------------------------------------ table


def test_a_table_downloads_as_shown_and_is_recorded(client, fx, team) -> None:  # noqa: F811
    from openpyxl import load_workbook

    body = {"what": "maintenance", "title": "Maintenance", "head": ["Address", "kWp", "Turned on"],
            "rows": [["7 Example Avenue", 16.465, "2024-09-01"], ["3 Trial Lane", None, None]]}  # fmt: skip
    r = client.post("/api/py/export/table.xlsx", headers=h(team["pm"]), json=body)
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content))["Maintenance"]
    assert [c.value for c in ws[1]] == ["Address", "kWp", "Turned on"]
    assert ws["A2"].value == "7 Example Avenue" and ws["B2"].value == 16.465 and ws["C2"].value.date().isoformat() == "2024-09-01"
    assert ws.max_row == 3
    last = fx.conn.execute("select what, row_count from data_exports where exported_by = %s order by export_id desc limit 1", (team["pm"]["uid"],)).fetchone()
    assert last == {"what": "maintenance", "row_count": 2}


@pytest.mark.parametrize(
    ("who", "what", "rows", "status"),
    [
        ("ho", "projects", [["x"]], 403),
        ("epc", "projects", [["x"]], 403),
        ("pm", "secrets", [["x"]], 400),
        ("pm", "projects", [["x", "y"]], 400),
        ("pm", "invoices", [["x"]], 403),  # Admin's list
    ],
)
def test_a_table_download_is_refused(client, team, who, what, rows, status) -> None:  # noqa: F811
    r = client.post("/api/py/export/table.xlsx", headers=h(team[who]), json={"what": what, "title": "T", "head": ["A"], "rows": rows})
    assert r.status_code == status, r.text
