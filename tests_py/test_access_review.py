"""The quarterly access review (ISO 27001 A.5.18), and the Solar lists' downloads (migration 0053)."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from _lib import clerk
from conftest import bearer


def test_a_superadmin_downloads_everyones_access(client, fx, monkeypatch) -> None:
    sa = fx.user("superadmin")
    pm = fx.user("project_manager")
    monkeypatch.setattr(
        clerk, "sign_in_facts", lambda: {pm["clerk_user_id"]: {"last_sign_in_ms": 1792650000000, "two_step": True}}
    )
    r = client.get("/api/py/export/access-review.xlsx", headers=bearer(sa["clerk_user_id"]))
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Access", "Sign-off"]
    rows = list(wb["Access"].iter_rows(values_only=True))
    head, mine = rows[0], next(x for x in rows if x[3] == pm["email"])
    assert head[:3] == ("S/N", "Name", "Role")
    assert mine[head.index("Two-step verification")] == "On" and mine[head.index("Last sign-in")] is not None
    assert [x[0] for x in wb["Sign-off"].iter_rows(values_only=True)][1] == "Reviewed by"


def test_only_a_superadmin(client, fx, monkeypatch) -> None:
    monkeypatch.setattr(clerk, "sign_in_facts", lambda: {})
    for who in ("project_manager", "finance", "homeowner"):
        u = fx.user(who)
        assert client.get("/api/py/export/access-review.xlsx", headers=bearer(u["clerk_user_id"])).status_code == 403


def test_without_clerk_it_still_downloads(client, fx, monkeypatch) -> None:
    def down():
        raise RuntimeError("no Clerk")

    monkeypatch.setattr(clerk, "sign_in_facts", down)
    sa = fx.user("superadmin")
    assert client.get("/api/py/export/access-review.xlsx", headers=bearer(sa["clerk_user_id"])).status_code == 200


def test_the_solar_lists_download(client, fx) -> None:
    pm = fx.user("project_manager")
    for what in ("solar", "solar_report"):
        r = client.post(
            "/api/py/export/table.xlsx",
            json={"what": what, "title": "Solar", "head": ["Plant", "kWh"], "rows": [["A plant", 12.5]]},
            headers=bearer(pm["clerk_user_id"]),
        )
        assert r.status_code == 200, r.text
